"""
DOCX generation engine for ExportManager.
"""

import re
import os
import json
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional
import logging

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .utils import (
    clean_markdown, topic_display_label, looks_like_json, 
    sanitize_json_content, download_image_to_temp, ensure_docx_compatible,
    clean_latex
)

logger = logging.getLogger(__name__)

class DocxEngine:
    def __init__(self, db):
        self.db = db

    def export_full_course(
        self,
        subject: str,
        all_notes_data: List[Dict],
        answer_key: Optional[str] = None,
        output_path: str = None,
        include_title_page: bool = True,
        chapter_summaries: Optional[Dict] = None,
        glossary: Optional[List[Dict]] = None
    ) -> str:
        doc = Document()
        
        if include_title_page:
            # 1. Premium Cover Page
            for _ in range(8): doc.add_paragraph()
            title = doc.add_heading(subject.upper(), 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
            sub = doc.add_paragraph()
            sub_run = sub.add_run("Comprehensive Study Guide & Instructor Manual")
            sub_run.italic = True
            sub_run.font.size = Pt(16)
            sub_run.font.color.rgb = RGBColor(0x2E, 0x59, 0x8A)
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for _ in range(10): doc.add_paragraph()
            gen_info = doc.add_paragraph()
            gen_info.add_run(f"Generated on {datetime.now().strftime('%B %d, %Y')}").italic = True
            gen_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_page_break()
            
            # 2. Table of Contents
            doc.add_heading("Table of Contents", 1)
            for idx, topic in enumerate(all_notes_data):
                t_id = topic.get('id', '')
                t_title = clean_markdown(topic.get('title', 'Untitled'))
                label = topic_display_label(t_id, t_title)
                depth = t_id.count('.')
                indent = "    " * depth
                bookmark_name = f"_topic_{idx}"
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(6 if depth == 0 else 1)
                hyperlink = OxmlElement('w:hyperlink')
                hyperlink.set(qn('w:anchor'), bookmark_name)
                run_el = OxmlElement('w:r')
                rPr = OxmlElement('w:rPr')
                rStyle = OxmlElement('w:rStyle')
                rStyle.set(qn('w:val'), 'Hyperlink')
                rPr.append(rStyle)
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), str((12 - depth) * 2))
                rPr.append(sz)
                if depth == 0:
                    b = OxmlElement('w:b')
                    rPr.append(b)
                    color_el = OxmlElement('w:color')
                    color_el.set(qn('w:val'), '1A5276')
                    rPr.append(color_el)
                else:
                    color_el = OxmlElement('w:color')
                    color_el.set(qn('w:val'), '333333')
                    rPr.append(color_el)
                run_el.append(rPr)
                text_el = OxmlElement('w:t')
                text_el.set(qn('xml:space'), 'preserve')
                text_el.text = f"{indent}{label}"
                run_el.append(text_el)
                hyperlink.append(run_el)
                p._p.append(hyperlink)

            doc.add_paragraph()
            hint = doc.add_paragraph()
            hint_run = hint.add_run("(Open in Word → right-click below → Update Field for page numbers)")
            hint_run.italic = True
            hint_run.font.size = Pt(8)
            hint_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            paragraph = doc.add_paragraph()
            run = paragraph.add_run()
            fldChar = OxmlElement('w:fldChar')
            fldChar.set(qn('w:fldCharType'), 'begin')
            run._r.append(fldChar)
            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
            run._r.append(instrText)
            fldChar = OxmlElement('w:fldChar')
            fldChar.set(qn('w:fldCharType'), 'separate')
            run._r.append(fldChar)
            fldChar = OxmlElement('w:fldChar')
            fldChar.set(qn('w:fldCharType'), 'end')
            run._r.append(fldChar)
            doc.add_page_break()
        
        # 3. Content Sections
        for i, topic in enumerate(all_notes_data):
            title = clean_markdown(topic.get('title', 'Untitled'))
            t_id = topic.get('id', '')
            depth = t_id.count('.')
            heading_level = min(depth + 1, 4)
            display = topic_display_label(t_id, title)
            if depth == 0 and i > 0:
                doc.add_page_break()
            heading_para = doc.add_heading(display, heading_level)
            bookmark_name = f"_topic_{i}"
            bookmark_start = OxmlElement('w:bookmarkStart')
            bookmark_start.set(qn('w:id'), str(i))
            bookmark_start.set(qn('w:name'), bookmark_name)
            heading_para._p.insert(0, bookmark_start)
            bookmark_end = OxmlElement('w:bookmarkEnd')
            bookmark_end.set(qn('w:id'), str(i))
            heading_para._p.append(bookmark_end)
            for el in topic.get("elements", []):
                self._add_element(doc, el, heading_offset=depth)
            
        # 4. Answer Key
        if answer_key:
            doc.add_page_break()
            doc.add_heading("Answer Key for Practice Problems", 1)
            self._add_markdown(doc, answer_key)

        # 5. Chapter Summaries
        if chapter_summaries:
            doc.add_page_break()
            doc.add_heading("Chapter Summaries", 1)
            for ch_id, ch_data in chapter_summaries.items():
                ch_title = ch_data.get("title", ch_id)
                doc.add_heading(clean_markdown(ch_title), 2)
                summary_text = ch_data.get("summary", "")
                if summary_text:
                    doc.add_paragraph(clean_markdown(summary_text))
                key_concepts = ch_data.get("key_concepts", [])
                for concept in key_concepts:
                    doc.add_paragraph(clean_markdown(str(concept)), style="List Bullet")
                outcomes = ch_data.get("learning_outcomes", [])
                for outcome in outcomes:
                    doc.add_paragraph(clean_markdown(str(outcome)), style="List Bullet")
                connection = ch_data.get("connections_to_next", "")
                if connection:
                    bridge = doc.add_paragraph()
                    bridge_run = bridge.add_run(f"→ {clean_markdown(connection)}")
                    bridge_run.italic = True
                    bridge_run.font.color.rgb = RGBColor(0x2E, 0x59, 0x8A)

        if output_path is None:
            ts = datetime.now().strftime("%H%M%S")
            safe_subject = re.sub(r'[^\w\s-]', '', subject)[:40].strip()
            filename = f"Export_{safe_subject}_v{ts}.docx"
            output_path = Path("outputs") / filename
            
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return str(output_path)

    def export_single_note(
        self,
        content_data: Dict,
        audience: Optional[str] = None,
        include_tags: bool = False,
        output_path: str = None
    ) -> str:
        doc = Document()
        doc.add_heading(f"{content_data['topic']}: {content_data['subtopic']}", 0)
        
        from .utils import remove_tags
        content = content_data['content']
        if not include_tags:
            content = remove_tags(content)
            
        doc.add_paragraph(f"Difficulty: {content_data['difficulty_level']}")
        if audience:
            doc.add_paragraph(f"Audience: {audience}")
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()
        
        # Simple markdown to docx for single notes
        self._add_markdown(doc, content)
        
        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_topic = re.sub(r'[^\w\s-]', '', content_data['topic'])[:30]
            filename = f"Export_{safe_topic}_{timestamp}.docx"
            output_path = Path("outputs") / filename
            
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return str(output_path)

    def _add_element(self, doc: Document, element: Dict, heading_offset: int = 0):
        etype = element.get("type")
        content = element.get("content", "")
        if etype == "text" and looks_like_json(content):
            content = sanitize_json_content(content)

        if etype == "heading":
            level = element.get("level", 2)
            doc.add_heading(clean_markdown(content), min(level + heading_offset + 1, 9))
        elif etype == "text":
            # Detect and handle markdown tables within text (improved regex)
            parts = re.split(r'((?:^|\n)\|.*\|(?:(?:\n\|.*\|)+|(?:\r\n\|.*\|)+))', content)
            for part in parts:
                if not part: continue
                if re.match(r'^\n?\|.*\|', part.strip()):
                    self._add_table_from_md(doc, part)
                elif part.strip():
                    self._add_text_with_lists(doc, part)
        elif etype == "code":
            clean_code = content.strip().replace('```', '')
            # Use a table for consistent background/border
            table = doc.add_table(rows=1, cols=1)
            table.style = 'Table Grid'
            cell = table.cell(0, 0)
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'F2F2F2')
            cell._tc.get_or_add_tcPr().append(shd)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(clean_code)
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            doc.add_paragraph() # Spacer
        elif etype == "equation":
            clean_eq = content.replace('$', '').strip()
            # Wrap in a subtle table/box
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Minimalist border would be nice, but 'Table Grid' is standard
            cell = table.cell(0, 0)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(clean_latex(clean_eq))
            run.italic = True
            run.font.size = Pt(12)
            doc.add_paragraph()
        elif etype == "callout":
            ctype = element.get("callout_type", "Note")
            color_hex = "F3F3F3"
            if "Takeaway" in ctype: color_hex = "D9EAD3"
            elif "Analogy" in ctype: color_hex = "FFF2CC"
            elif "Understanding" in ctype: color_hex = "DEEBF7"
            elif "Practice" in ctype: color_hex = "FAD7A0"
            table = doc.add_table(rows=1, cols=1)
            table.style = 'Table Grid'
            cell = table.cell(0, 0)
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), color_hex)
            cell._tc.get_or_add_tcPr().append(shd)
            p = cell.paragraphs[0]
            header = p.add_run(f"{ctype}\n")
            header.bold = True
            header.font.size = Pt(11)
            p.add_run(clean_markdown(content))
            doc.add_paragraph()
        elif etype == "image_placeholder":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            local_img = element.get("image_local_path", "")
            image_url = element.get("image_url")
            temp_img_path = None
            is_temp = False
            if local_img and os.path.isfile(local_img):
                temp_img_path = local_img
            elif image_url:
                temp_img_path = download_image_to_temp(image_url)
                is_temp = True
            if temp_img_path:
                converted_path = ensure_docx_compatible(temp_img_path)
                try:
                    run = p.add_run()
                    run.add_picture(converted_path, width=Inches(5.5))
                    doc.add_paragraph()
                    if element.get("caption"):
                        cap_p = doc.add_paragraph()
                        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap_run = cap_p.add_run(f"Figure: {element.get('caption')}")
                        cap_run.italic = True
                        cap_run.font.size = Pt(9)
                        cap_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                except Exception as e:
                    logger.error(f"Failed to insert image into DOCX: {e}")
                finally:
                    if converted_path != temp_img_path:
                        try: os.remove(converted_path)
                        except: pass
                    if is_temp and temp_img_path:
                        try: os.remove(temp_img_path)
                        except: pass
        elif etype == "source_citation":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(f"Source: {clean_markdown(content)}")
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif etype == "table":
            rows = content.split('\n')
            grid = []
            for r in rows:
                if '---' in r: continue
                cells = [c.strip() for c in r.split('|') if c.strip() or r.startswith('|')]
                if any(cells): grid.append(cells)
            if not grid: return
            table = doc.add_table(rows=len(grid), cols=len(grid[0]))
            table.style = 'Table Grid'
            for r_idx, row_data in enumerate(grid):
                for c_idx, cell_data in enumerate(row_data):
                    if c_idx < len(table.columns):
                        table.cell(r_idx, c_idx).text = clean_latex(cell_data)

    def _add_table_from_md(self, doc, md_table):
        rows = md_table.strip().split('\n')
        grid = []
        for r in rows:
            if not r or re.match(r'^\|?\s*:?-+', r): continue
            cells = [c.strip() for c in r.split('|') if c.strip() or r.startswith('|')]
            if any(cells): grid.append(cells)
        if not grid: return
        
        table = doc.add_table(rows=len(grid), cols=len(grid[0]))
        table.style = 'Table Grid'
        for r_idx, row_data in enumerate(grid):
            for c_idx, cell_data in enumerate(row_data):
                if c_idx < len(table.columns):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = clean_latex(cell_data)
                    if r_idx == 0:
                        # Make header bold
                        for p in cell.paragraphs:
                            for r in p.runs: r.bold = True

    def _add_text_with_lists(self, doc, content):
        lines = content.split('\n')
        current_p = []
        def _flush():
            if current_p:
                doc.add_paragraph(clean_markdown("\n".join(current_p)))
                current_p.clear()
        for line in lines:
            if re.match(r"^\s*([-\*\+•]\s+|\d+[\.\\)]\s+)", line):
                _flush()
                clean = re.sub(r"^\s*([-\*\+•]\s*|\d+[\.\\)]\s*)", "", line.strip(), 1)
                doc.add_paragraph(clean_markdown(clean), style='List Bullet')
            else:
                current_p.append(line)
        _flush()

    def _add_markdown(self, doc: Document, markdown_text: str):
        lines = markdown_text.split('\n')
        current_p: list[str] = []
        def _flush():
            if current_p:
                text = "\n".join(current_p)
                p = doc.add_paragraph()
                self._add_rich_runs(p, text)
                current_p.clear()
        for line in lines:
            stripped = line.strip()
            if not stripped:
                _flush()
                continue
            hm = re.match(r'^(#{1,6})\s+(.*)', stripped)
            if hm:
                _flush()
                level = min(len(hm.group(1)) + 1, 9)
                doc.add_heading(clean_markdown(hm.group(2)), level)
                continue
            if re.match(r'^[-*_]{3,}$', stripped):
                _flush()
                continue
            bm = re.match(r'^([-*+])\s+(.*)', stripped)
            if bm:
                _flush()
                p = doc.add_paragraph(style='List Bullet')
                self._add_rich_runs(p, bm.group(2))
                continue
            nm = re.match(r'^(\d+)[.)]\s+(.*)', stripped)
            if nm:
                _flush()
                p = doc.add_paragraph(style='List Number')
                self._add_rich_runs(p, nm.group(2))
                continue
            current_p.append(line)
        _flush()

    def _add_rich_runs(self, paragraph, text: str):
        text = clean_latex(text)
        parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
        for part in parts:
            if not part: continue
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif part.startswith('`') and part.endswith('`'):
                run = paragraph.add_run(part[1:-1])
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
            else:
                clean = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', part)
                paragraph.add_run(clean)
